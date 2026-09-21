# -*- coding: utf-8 -*-
import os
_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_CW = os.path.join(_ROOT, "docs", "crosswalk")
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
CLUSTER_FILL = PatternFill("solid", fgColor="D9E2F3")
FIT = {
    "Direct":     PatternFill("solid", fgColor="D6EFD8"),
    "Partial":    PatternFill("solid", fgColor="FFF2CC"),
    "Properties": PatternFill("solid", fgColor="FCE4D6"),
    "None":       PatternFill("solid", fgColor="F8CBCB"),
}
THIN = Side(style="thin", color="BFBFBF")
BORD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()

def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = HDR_FILL
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        cell.border = BORD
    ws.row_dimensions[row].height = 34

def finish(ws, widths, ncols, freeze="A2"):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{ws.max_row}"

# ============================================================
# SHEET 1 - the 50-element AI SBOM crosswalk
# ============================================================
ws = wb.active
ws.title = "AI SBOM crosswalk"

HEAD = ["#", "Cluster", "Element", "What it captures", "Stated format / constraint",
        "CycloneDX 1.7 path", "CDX fit", "SPDX 3.0.1 path", "SPDX fit", "Gap or trap"]
ws.append(HEAD)
style_header(ws, len(HEAD))

R = []  # cluster, element, meaning, constraint, cdx, cdxfit, spdx, spdxfit, note

M = "Metadata"
R += [
(M,"SBOM author","Entity that generated the SBOM. Distinct from the producer of the thing described.","String. Full names, no acronyms unless official.","metadata.authors[] (organizationalContact); metadata.manufacturer","Direct","CreationInfo.createdBy -> Agent","Direct",""),
(M,"SBOM version","Version of the SBOM document itself, per component name/version pair.","SemVer 2.0.0 (major SHOULD be 1) or RFC 9562 UUID.","version (integer at BOM root)","Partial","No document-version property","None","CDX version is an integer, not SemVer. SPDX 3.0.1 has no SBOM document version at all. Real gap in both."),
(M,"SBOM data format name","Which format this SBOM is written in.","Free text.","bomFormat (const 'CycloneDX')","Direct","@context + type: SpdxDocument","Direct",""),
(M,"SBOM data format version","Version of that format. Deprecated versions should not be used.","Free text. Deprecation rule.","specVersion (e.g. '1.7')","Direct","CreationInfo.specVersion (e.g. '3.0.1')","Direct",""),
(M,"SBOM author signature","Digital signature of the SBOM author, for integrity and authenticity.","Algorithm approved by NIST DSS, ISO/IEC 14888-4:2024 or ENISA. No format named.","signature (JSF, jsf-0.82 schema); also declarations[].signature and citation[].signature","Direct","No native signature construct","None","THE COMMERCIAL HOOK. CycloneDX has a signature slot; SPDX 3.0.1 does not, so SPDX signing is out of band. Neither spec names a signature format (no Sigstore, DSSE, JWS or in-toto)."),
(M,"SBOM tool name","Tool the author used to build the SBOM.","String, full names.","metadata.tools.components[].name","Direct","CreationInfo.createdUsing[] -> Tool.name","Direct",""),
(M,"SBOM tool version","Version of that tool. If none, say 'unknown'.","String. Explicit 'unknown' required if absent.","metadata.tools.components[].version","Direct","Tool as Package -> packageVersion","Partial","Empty is non-conformant. Explicit 'unknown' is conformant. Neither format enforces this."),
(M,"SBOM generation context","Lifecycle phase at which the SBOM was made.","Open list. Examples: before build, build, after build.","metadata.lifecycles[].phase (design, pre-build, build, post-build, operations, discovery, decommission)","Partial","LifecycleScopeType on relationships (design, build, development, test, runtime, other)","Partial","Vocabularies do not line up with each other or with the spec's three example values. SPDX scopes relationships, not the document."),
(M,"SBOM timestamp","When the SBOM data was last changed. Every version gets a new one.","RFC 9557 (extended timestamps with bracketed suffixes).","metadata.timestamp (JSON Schema format: date-time)","Partial","CreationInfo.created (xsd:dateTimeStamp)","Partial","RFC 9557 permits suffixes such as [America/Chicago] that RFC 3339 validators reject. Both formats validate as plain date-time, so a strictly conformant timestamp can fail their schemas."),
(M,"SBOM dependency relationship","Component X includes Y; or Y is derived from or a descendant of Z.","Values: includes, included in, plus derived-from and descendant-of semantics.","dependencies[] (ref, dependsOn, provides); pedigree.ancestors/descendants/variants for derivation","Direct","Relationship: contains, dependsOn, ancestorOf, descendantOf, hasVariant","Direct",""),
]

S = "System Level Properties"
R += [
(S,"System name","Human-readable name of the whole AI system.","Multi-valued: alternate names should be capturable.","metadata.component.name","Partial","Package.name","Partial","Neither format allows alternate names on a single component. Spec explicitly asks for multiple entries. Use properties or externalIdentifier as a workaround."),
(S,"System components","Everything included in the AI system: models, databases, tools.","Free.","metadata.component.components[]; components[] + dependencies[]","Direct","Relationship contains","Direct",""),
(S,"System producer","Entity that creates, defines and identifies the system.","Free.","metadata.component.manufacturer / .supplier","Direct","Artifact.originatedBy / suppliedBy","Direct",""),
(S,"System version","Version of the system. If none, say 'unknown'.","Explicit 'unknown' required if absent.","metadata.component.version","Direct","Package.packageVersion","Direct",""),
(S,"System timestamp","Last update to the system, including model or software updates.","No RFC constraint stated (unlike SBOM timestamp).","No per-component update field. component.releaseNotes.timestamp is the nearest.","Partial","Artifact.builtTime / releaseTime / validUntilTime","Direct","SPDX is materially better here. CycloneDX has no component-level 'last updated' field."),
(S,"System data flow","Data movement between components: endpoints, APIs, agent protocols, external services.","Link, reference or description.","services[].endpoints[]; services[].data[] with flow (inbound, outbound, bi-directional, unknown) and classification","Direct","No data-flow construct","None","CycloneDX SaaSBOM services model this properly. SPDX 3.0.1 cannot express data flow at all. Biggest single divergence between the two formats."),
(S,"System data usage","How data is processed and consumed, including logging and derived metadata.","Link to documentation.","services[].data[].classification; externalReferences[]","Partial","AIPackage.informationAboutApplication (free text)","Partial","Both weak. No structured notion of 'is user input used for training'."),
(S,"System input/output properties","Input and output types, modality, preprocessing, tokenizer.","Free.","modelCard.modelParameters.inputs[].format / outputs[].format","Partial","AIPackage.informationAboutApplication (free text)","Partial","CycloneDX allows exactly ONE key here (format) because inputOutputMLParameters is additionalProperties:false. Modality lives in the separate cdx:ai-ml property taxonomy, not the schema. This is model-level in CDX, not system-level."),
(S,"Intended application area","Domain the system is deployed in: healthcare, finance, cybersecurity, real time.","Free.","modelCard.considerations.useCases[]","Partial","AIPackage.domain (0..*)","Direct","SPDX has a purpose-built field. CycloneDX only has model-level use cases."),
]

MO = "Models"
R += [
(MO,"Model name","Human-readable model name.","Multi-valued: alternate names should be capturable.","components[].name (type: machine-learning-model)","Partial","AIPackage.name","Partial","Same alternate-names gap as System name."),
(MO,"Model identifier","One or more machine-readable identifiers.","CPE, PURL, UUID, org IDs, commit hashes, OmniBOR, SWHID. Include ALL of them.","components[].purl, .cpe, .swid, .omniborId, .swhid, .bom-ref","Direct","Element.spdxId; Element.externalIdentifier[]; Package.packageUrl","Direct","Both support the named schemes. Hugging Face PURL type (pkg:huggingface/...) is the de facto model identifier."),
(MO,"Model version","Model version. If none, say 'unknown'.","Explicit 'unknown' required if absent.","components[].version","Direct","Package.packageVersion","Direct",""),
(MO,"Model timestamp","Last update to the model, or its production release date.","Free.","components[].releaseNotes.timestamp","Partial","Artifact.releaseTime / builtTime","Direct","SPDX better."),
(MO,"Model producer","Who pre-trained, post-trained or fine-tuned the model. Also the security point of contact.","Multi-valued.","components[].manufacturer, .supplier, .authors[], .publisher","Direct","Artifact.originatedBy (0..*), suppliedBy (0..1)","Direct",""),
(MO,"Model description","Capabilities, known limitations, and lineage including the predecessor model.","Free.","components[].description; modelCard.considerations.technicalLimitations[]; pedigree.ancestors[]","Partial","Element.description; AIPackage.limitation; AIPackage.informationAboutTraining; ancestorOf relationship","Partial","LINEAGE IS THE PROBLEM. CycloneDX pedigree is structured but generic. SPDX 3.0.1 has no base-model edge at all, and even 3.1-dev's finetunedOn points at a dataset, not a parent model. 'This was fine-tuned from Llama 3' has no interoperable representation in either format."),
(MO,"Model hash value","Hash of the weights or model file.","ASCII-formatted value.","components[].hashes[].content","Direct","Element.verifiedUsing -> Hash.hashValue","Direct",""),
(MO,"Model hash algorithm","Algorithm that produced the hash.","IANA Hash Function Textual Names. NIST-approved.","components[].hashes[].alg (closed enum: MD5, SHA-1, SHA-256/384/512, SHA3-*, BLAKE2b-*, BLAKE3)","Partial","Hash.algorithm (SPDX HashAlgorithm vocabulary)","Partial","NEITHER FORMAT USES IANA NAMES. Both have their own closed enums. A conformance checker has to map IANA names onto each format's vocabulary, and the sets do not fully overlap."),
(MO,"Model properties","Architecture, parameter count, network type, hyperparameters. Explicitly nestable.","May contain sub-elements.","modelCard.modelParameters.approach.type (enum: supervised, unsupervised, reinforcement-learning, semi-supervised, self-supervised), .task, .architectureFamily, .modelArchitecture; plus cdx:ai-ml:model:parameter:count and :hyperparameter:* taxonomy","Direct","AIPackage.typeOfModel (0..*); AIPackage.hyperparameter[] (DictionaryEntry key/value strings)","Direct","CycloneDX is richer via the property taxonomy but the taxonomy sits outside the schema. SPDX hyperparameters are untyped strings with no units."),
(MO,"Model input-output properties","Input and output types, context length, modality, tokenizer.","Free.","modelCard.modelParameters.inputs[].format / outputs[].format only","Partial","No dedicated property","None","Context length and modality have no schema home in either format. CycloneDX puts modality in the taxonomy (cdx:ai-ml:model:modality); context length only as a hyperparameter property."),
(MO,"Model training properties","All training: pre-training, fine-tuning, continual learning, RLHF, DPO, PPO, GRPO.","Free, link to model card.","modelCard.modelParameters.approach.type; formulation[].workflows[] for the training run; cdx:ai-ml:model:parameter:tune_method (full, sft, rlhf, adapter, prompt, lora, alora, qlora)","Partial","AIPackage.informationAboutTraining (free text); trainedOn relationship","Partial","CycloneDX formulation can model the actual training run as a workflow. SPDX models it as a Build. Both are possible, neither is conventional, and no two producers will do it the same way."),
(MO,"Model license","Licence type, and whether open weight, open architecture, open data or open training.","Link to licence doc, or the SPDX/CDX licence fields.","components[].licenses[] (SPDX expression)","Partial","hasDeclaredLicense and hasConcludedLicense relationships (exactly one of each REQUIRED on an AIPackage)","Partial","OPEN-WEIGHT VOCABULARY IS MISSING EVERYWHERE. Neither format has any way to say open weight vs open architecture vs open data vs open training. This is a named requirement with no representation."),
(MO,"Model external references","Links to model cards, system cards, papers, documentation.","Link/URL.","components[].externalReferences[] (47-value type enum incl. model-card, documentation, attestation, threat-model, adversary-model, risk-assessment, certification-report, pentest-report)","Direct","Element.externalRef[]","Direct","CycloneDX has a dedicated model-card reference type. SPDX external refs are less typed."),
]

D = "Datasets Properties"
R += [
(D,"Dataset name","Name given by the dataset creator.","Free.","componentData.name; or components[].name with type: data","Direct","DatasetPackage.name","Direct",""),
(D,"Dataset description","What the dataset is for: pre-training, fine-tuning, benchmark, evaluation. Public or private.","Free.","componentData.description","Direct","Element.description; DatasetPackage.intendedUse","Direct",""),
(D,"Dataset content","What is in it and in what shape: financial, medical, image, audio, JSON, XML.","Free.","componentData.type (enum: source-code, configuration, dataset, definition, other); .contents.{attachment,url,properties}","Partial","DatasetPackage.datasetType (14-value enum: audio, categorical, graph, image, noAssertion, numeric, other, sensor, structured, syntactic, text, timeseries, timestamp, video); .datasetSize (bytes)","Direct","SPDX is clearly better. Its 14-value datasetType is the only real content vocabulary in either format, and it is the ONLY required property in the whole SPDX AI/Dataset surface (minCount 1)."),
(D,"Dataset identifier","Unique identification of the dataset.","URL, URI.","componentData.bom-ref; components[].purl","Direct","Element.spdxId; externalIdentifier[]; Package.downloadLocation","Direct",""),
(D,"Dataset hash","Hash over the dataset file.","No algorithm field is defined (unlike Models).","components[].hashes[] on the wrapping component only","Partial","Element.verifiedUsing -> Hash","Direct","TRAP: CycloneDX componentData has NO hash property of its own. You can only hash the wrapping component. Also note the spec itself omits a Dataset hash algorithm element, which is an asymmetry with the Models cluster."),
(D,"Dataset provenance","Origin, collection method (web crawl, commercial agreement), post-processing, curation, labelling, creator, synthetic-data method.","Free.","componentData.governance.{custodians,stewards,owners}; .contents.url; formulation[] for the pipeline","Partial","DatasetPackage.dataCollectionProcess; .dataPreprocessing (0..*); .anonymizationMethodUsed; Artifact.originatedBy","Direct","SPDX has purpose-built fields for collection and preprocessing. CycloneDX has stewardship roles SPDX lacks. Neither has a synthetic-data-method field."),
(D,"Dataset statistical properties","Mean, variance, median, mode, range, skewness across the dataset lifecycle.","Free.","componentData.properties[] or .graphics","Properties","No property","None","REAL GAP IN BOTH. No statistical characterisation exists anywhere in either format. Nearest is SPDX datasetNoise (free text) and CycloneDX graphics (base64 images)."),
(D,"Dataset sensitivity","Whether it contains PII, freely accessible, copyright-protected, sensitive (financial, medical) or national-security data.","Free.","componentData.sensitiveData[] (free-text array); .classification (unconstrained string)","Partial","DatasetPackage.hasSensitivePersonalInformation (PresenceType: yes/no/noAssertion); .confidentialityLevel (TLP: red, amber, green, clear)","Partial","SPDX has controlled vocabularies, CycloneDX has none. But NEITHER covers copyright-protected or national-security categories, which the spec names explicitly. SPDX 3.1-dev adds amberStrict."),
(D,"Dataset dependency relationship","Software used to create, modify and maintain the dataset. Derivation from another dataset.","Free.","dependencies[] on the data component; formulation[].components[] for the tooling","Partial","Relationship: dependsOn, usesTool, ancestorOf; Build profile for the pipeline","Partial","Both express it, neither conventionally."),
(D,"Dataset license","Licence type of the dataset.","Link to licence doc.","components[].licenses[]","Direct","hasDeclaredLicense and hasConcludedLicense (exactly one each REQUIRED on a DatasetPackage)","Direct",""),
]

I = "Infrastructure"
R += [
(I,"Infrastructure software","Firmware, package managers, third-party libraries, frameworks, runtimes, tools needed to run the system.","Free.","components[] of type library/framework/container/operating-system + dependencies[]; formulation[].components[]","Direct","Relationship dependsOn / usesTool; Build profile","Direct",""),
(I,"Infrastructure hardware","Link to an existing HBOM covering the hardware the system runs on.","Link to an HBOM.","externalReferences[].type: 'bom' with a BOM-Link URN; components[].type: 'device'","Partial","No Hardware profile in 3.0.1 (arrives in 3.1-dev with runsOn)","None","SPDX 3.0.1 cannot express hardware at all. CycloneDX can link a separate BOM but has no HBOM profile either. Note OMB M-26-05 now makes agencies inventory hardware as well as software, so this gap has a buyer."),
]

SP = "Security Properties"
R += [
(SP,"Security controls","General controls (encryption, data minimisation, differential privacy, access control, API auth, anomaly detection) AND AI-specific controls (adversarial robustness training, prompt injection controls, I/O filters, training-data curation).","Free, may link to a framework.","properties[] under a private namespace; externalReferences[].type: threat-model / adversary-model / risk-assessment; declarations[] (CDXA) against definitions.standards[]","Properties","Extension profile; AIPackage.standardCompliance (partial)","Properties","THE BIGGEST GAP ON THIS SHEET. Neither format has ANY field for adversarial robustness, prompt injection resistance, jailbreak resistance, data poisoning, watermarking, guardrails or red-teaming. Zero hits for all of those terms in both schemas. This is unclaimed territory."),
(SP,"Security compliance","Certification schemes, standards or frameworks the model or system complies with.","Free.","declarations[] (CDXA); externalReferences[].type: certification-report","Partial","AIPackage.standardCompliance (xsd:string, 0..*)","Direct","SPDX has a plain field. CycloneDX requires the heavier Attestations construct."),
(SP,"Cybersecurity policy information","Link to the producer's published security.txt.","URL to a security.txt file.","externalReferences[] with an appropriate type","Partial","Element.externalRef[] (securityAdvisory, securityOther)","Partial","Neither has a security.txt reference type by name. Nearest fit only."),
(SP,"Vulnerability referencing","Link to databases or repositories showing exploitability of known vulnerabilities.","Static link/URL.","vulnerabilities[] (CVE, CWE, CVSS ratings, VEX analysis state); externalReferences[].type: vulnerability-assertion","Direct","Security profile: Vulnerability, VulnAssessmentRelationship, VEX relationships, CVSS v2/v3/v4, EPSS, SSVC","Direct","Both strong, but both are CVE-shaped. Neither can express a model-level attack surface, only vulnerabilities in software dependencies."),
]

K = "Key Performance Indicators"
R += [
(K,"Security metrics","Metrics for the security characteristics of the models or the system, such as robustness benchmarks.","Free.","modelCard.quantitativeAnalysis.performanceMetrics[] (type free text, value is a STRING, optional slice and confidence interval)","Properties","AIPackage.metric[] (DictionaryEntry: string key, string value)","Properties","Untyped in both. No units, no thresholds beyond SPDX metricDecisionThreshold, no security semantics. And in CycloneDX there is NO field linking a metric to the dataset it was measured on."),
(K,"Operational performance KPIs","Uptime, incident resolution time, latency, request throughput, load balancing.","Free.","properties[] only","Properties","AIPackage.metric[] only","Properties","REAL GAP IN BOTH. No operational-telemetry concept exists in either format. This is arguably out of scope for a static BOM, and worth saying so."),
]

for idx, (cl, el, mean, cons, cdx, cdxf, spdx, spdxf, note) in enumerate(R, start=1):
    ws.append([idx, cl, el, mean, cons, cdx, cdxf, spdx, spdxf, note])
    r = ws.max_row
    for c in range(1, len(HEAD) + 1):
        cell = ws.cell(row=r, column=c)
        cell.font = Font(name=FONT, size=9)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = BORD
    ws.cell(row=r, column=1).alignment = Alignment(vertical="top", horizontal="center")
    ws.cell(row=r, column=2).fill = CLUSTER_FILL
    ws.cell(row=r, column=3).font = Font(name=FONT, size=9, bold=True)
    for col, fit in ((7, cdxf), (9, spdxf)):
        cell = ws.cell(row=r, column=col)
        cell.fill = FIT[fit]
        cell.alignment = Alignment(vertical="top", horizontal="center", wrap_text=True)
        cell.font = Font(name=FONT, size=9, bold=True)

finish(ws, [4, 17, 25, 40, 26, 46, 9, 42, 9, 60], len(HEAD))
print("sheet1 rows:", ws.max_row - 1)

# ============================================================
# SHEET 2 - the 17 general SBOM fields (v2.1, 29 July 2026)
# ============================================================
ws2 = wb.create_sheet("SBOM v2.1 - 17 fields")
H2 = ["#", "Category", "Field", "Change vs 2021", "Definition (verbatim)", "Type / rule",
      "CycloneDX 1.7 path", "SPDX 3.0.1 path", "Note"]
ws2.append(H2)
style_header(ws2, len(H2))

F = []
MD = "SBOM Metadata"
F += [
(MD,"SBOM Author","Major update","The name of the entity that creates the SBOM data for the target component.","String. Full names, no acronyms unless official. Distinct from Component Producer.","metadata.authors[]","CreationInfo.createdBy","Same wording as the AI SBOM metadata element."),
(MD,"SBOM Author Signature","NEW","A digital signature attributable to the SBOM author.","Algorithm per NIST DSS, ISO/IEC 14888-4:2024 or ENISA Agreed Cryptographic Mechanisms. No signature FORMAT named.","signature (jsf-0.82)","None native","Attests integrity and authenticity of the document ONLY. Accuracy, coverage and completeness are explicitly out of scope. The word 'attestation' does not appear in the document."),
(MD,"SBOM Data Format Name","NEW","The name of the data format used to represent the SBOM data.","Free text; defers to Machine-Processable Data.","bomFormat","@context / type","No enumeration of allowed values."),
(MD,"SBOM Data Format Version","NEW","Identifier designated by the SBOM data format to specify the version of the data format.","Deprecated versions should not be used. No version numbers named.","specVersion","CreationInfo.specVersion","The only version rule is negative: do not use deprecated versions."),
(MD,"SBOM Generation Context","NEW","The relative software lifecycle phase and data available at the time the SBOM author generated the SBOM.","OPEN list. Examples: before build, build, after build. More specific identifiers allowed.","metadata.lifecycles[].phase","LifecycleScopeType","Do not hard-reject unknown values. Note the coupling: a 'before build' SBOM legitimately has Component Hash Value = unknown."),
(MD,"SBOM Timestamp","Minor update","Record of the date and time of the most recent update to the SBOM data.","RFC 9557 (NOT RFC 3339). Each version gets a new timestamp.","metadata.timestamp","CreationInfo.created","A validator that only accepts RFC 3339 will reject conformant timestamps."),
(MD,"SBOM Tool Name","NEW","The name of the tool used by the SBOM author to generate or amend the SBOM.","String, full names.","metadata.tools.components[].name","CreationInfo.createdUsing[]",""),
(MD,"SBOM Tool Version","NEW","Identifier for the version of the tool identified in the SBOM Tool Name element.","If none available, state 'unknown'.","metadata.tools.components[].version","Tool packageVersion","Empty is non-conformant; explicit 'unknown' is conformant."),
(MD,"SBOM Version","NEW","Identifier designated by the SBOM author to specify a change in the SBOM document from a previously identified version.","SemVer 2.0.0 (major SHOULD be '1' to signal conformance to these minimum elements) or RFC 9562 UUID.","version (integer)","None","Major version '1' is a directly checkable conformance signal. CycloneDX version is an integer, so SemVer does not fit."),
]
CD = "Component Data"
F += [
(CD,"Component Producer","Major update","The name of an entity that creates, defines, and identifies components.","EXACTLY ONE per component. Fallback: explicitly state 'unknown provenance'.","components[].manufacturer / .supplier","Artifact.originatedBy","Replaces 2021 'Supplier Name', which was ambiguous around distributors. Footnote warns 'producer' is not the EU term 'manufacturer'."),
(CD,"Component Dependency Relationship","Minor update","The relationship between two components, where one component is necessary for the operation of the other.","Embedded or linked SBOMs both acceptable.","dependencies[]","Relationship contains / dependsOn","Copy-pasted code should be documented as an inclusion dependency named as if it were a local fork."),
(CD,"Component Hash Value","NEW","The output generated from applying a cryptographic hash algorithm to an executable component artifact.","ASCII, hexadecimal encoded. If no access to the artifact, state 'unknown'.","components[].hashes[].content","Element.verifiedUsing -> Hash.hashValue",""),
(CD,"Component Hash Algorithm","NEW","The cryptographic algorithm used to compute the Component Hash Value.","IANA Hash Function Textual Names. NIST-approved algorithm.","components[].hashes[].alg (own closed enum)","Hash.algorithm (own vocabulary)","Neither format uses IANA names. You must map."),
(CD,"Component Identifiers","Major update","Identifiers used to identify a component or serve as a look-up key for relevant databases.","At least ONE. Include ALL known. CPE, PURL (now ECMA-427), UUID, org IDs, commit hashes, OmniBOR, SWHID (now ISO/IEC 18670:2025).","components[].purl / .cpe / .swid / .omniborId / .swhid","externalIdentifier[]; packageUrl",""),
(CD,"Component License","NEW","The identifier(s) for the license(s) under which the software component is available.","Prefer SPDX licence identifiers. URL fallback. Must flag proprietary conditions. Explicit 'unknown' if not known.","components[].licenses[]","hasDeclaredLicense / hasConcludedLicense","Definition differs slightly between the body text and Appendix A Table 1."),
(CD,"Component Name","Minor update","The name assigned by the component producer to a software component.","Multiple entries allowed (this is the 2026 change).","components[].name","Element.name","Neither format actually supports alternate names on one component."),
(CD,"Component Version","Major update","Identifier used by the component producer to specify a change in a software component.","If the producer gives none, state 'unknown'.","components[].version","Package.packageVersion",""),
]

for idx, row in enumerate(F, start=1):
    ws2.append([idx] + list(row))
    r = ws2.max_row
    for c in range(1, len(H2) + 1):
        cell = ws2.cell(row=r, column=c)
        cell.font = Font(name=FONT, size=9)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = BORD
    ws2.cell(row=r, column=1).alignment = Alignment(vertical="top", horizontal="center")
    ws2.cell(row=r, column=2).fill = CLUSTER_FILL
    ws2.cell(row=r, column=3).font = Font(name=FONT, size=9, bold=True)
    if row[2] == "NEW":
        ws2.cell(row=r, column=4).fill = PatternFill("solid", fgColor="D6EFD8")
    ws2.cell(row=r, column=4).alignment = Alignment(vertical="top", horizontal="center", wrap_text=True)
finish(ws2, [4, 17, 30, 14, 56, 44, 34, 32, 58], len(H2))

# ============================================================
# SHEET 3 - the six practices
# ============================================================
ws3 = wb.create_sheet("Practices (6)")
H3 = ["#", "Practice", "Change", "What it requires", "What it means for a checker"]
ws3.append(H3)
style_header(ws3, len(H3))
P = [
("Accommodation of Updates to SBOM Data","Major","Organizations should accommodate updates including corrections. SBOM authors should correct errors promptly. Errors may be weighed in risk decisions.","Replaces 2021 'Accommodation of Mistakes'. The stated reason is that data quality has improved enough that recipients can now expect accuracy."),
("Coverage","Major","All components including transitive dependencies. THERE IS NO MINIMUM DEPTH. Multiple instances listed separately. May exclude non-code files but may include security-relevant ones such as config files. Linking to other SBOMs is allowed provided the recipient can access all of them.","Replaces 2021 'Depth', which required only top-level dependencies. 'No minimum depth' means no threshold is set BECAUSE full coverage is expected, not that shallow is fine. The completeness contract: absence of a component must be meaningful enough that a recipient can conclude a new CVE does not affect them."),
("Distribution and Delivery","Minor","Available promptly to those who need them. Access controls may limit unauthorized sharing but must not block authorized parties or prevent integration into trusted security tools. Mechanisms: accompanying installation, a version-specific URL, an API to a database, a public repository.","The standalone 2021 'Access Control' element was REMOVED and folded in here. A version-specific URL and an API are both explicitly blessed, which is the opening for a hosted service."),
("Explicitly Identifying Unknown Information","Major","If a field is not provided, state whether it is UNKNOWN to the author or WITHHELD by the author. The author should have a process for recipients to ask about redacted security-related information. An SBOM may be considered incomplete if essential data is withheld.","THREE STATES, not two: present, unknown, withheld. A checker must model all three. Silence or empty is non-conformant for SBOM Tool Version, Component Producer, Component Hash Value, Component License and Component Version."),
("Frequency","Minor","Each software version or update gets an SBOM. A new build or release triggers a new SBOM, including builds that integrate updated components. New details or a corrected error trigger a revised SBOM.","Each new SBOM gets a new timestamp and an incremented SBOM Version, scoped per component name/version pair. The SBOM version axis is independent of the component version axis."),
("Machine-Processable Data","Major","SPDX (ISO/IEC 5962:2021) and CycloneDX (ECMA-424) are the two named formats. Support all widely used, open source, compatible formats. Reassess regularly. Avoid deprecated versions.","SWID Tags were REMOVED from the accepted format list. No preference is expressed between SPDX and CycloneDX, and NO version numbers are named for either."),
]
for idx, row in enumerate(P, start=1):
    ws3.append([idx] + list(row))
    r = ws3.max_row
    for c in range(1, len(H3) + 1):
        cell = ws3.cell(row=r, column=c)
        cell.font = Font(name=FONT, size=9)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = BORD
    ws3.cell(row=r, column=1).alignment = Alignment(vertical="top", horizontal="center")
    ws3.cell(row=r, column=2).font = Font(name=FONT, size=9, bold=True)
    ws3.cell(row=r, column=3).alignment = Alignment(vertical="top", horizontal="center")
finish(ws3, [4, 40, 10, 78, 88], len(H3))
print("sheets 2-3 done")

# ============================================================
# SHEET 4 - traps and open gaps
# ============================================================
ws4 = wb.create_sheet("Traps and gaps")
H4 = ["#", "Where", "What bites you", "Why it matters to the product"]
ws4.append(H4)
style_header(ws4, len(H4))
T = [
("AI SBOM spec","No element is marked required, optional or recommended. There is no MUST/SHOULD/MAY anywhere, and the whole set is declared 'not mandatory'.","You have to invent the conformance policy. That is the product. Ship a documented default profile and let people override it."),
("AI SBOM spec","44 of the 50 elements have NO data type. Only SBOM timestamp (RFC 9557), SBOM version (SemVer or UUID), Model hash value (ASCII), Model hash algorithm (IANA names), Model identifier (CPE/PURL/UUID/OmniBOR/SWHID) and SBOM generation context carry any constraint.","Everything else is free text or a link. A checker can only verify presence, not correctness, for most of the surface."),
("AI SBOM spec","There is NO format binding. 'CycloneDX' does not appear in the document at all. 'SPDX' appears once, in an example.","No official crosswalk exists. This spreadsheet is the thing that does not exist yet."),
("AI SBOM spec","No section on sharing, distribution, access control or confidentiality. The 2021 NTIA elements had a dedicated Access Control practice; this has nothing.","Notable omission given Dataset sensitivity can flag PII and national-security data. Nothing says how that affects disclosure."),
("AI SBOM spec","No appendices, no glossary, no worked example, no JSON sample, no conformance checklist, no crosswalk. 24 pages, one figure, seven tables.","There is no reference implementation and no test data. Both are yours to build."),
("AI SBOM spec","Naming inconsistencies to normalise deliberately: 'Model input-output properties' (hyphen) vs 'System input/output properties' (slash). Section heading 'Datasets Properties' vs field prefix 'Dataset ...'. Element 'Model version' vs prose 'AI model version element'.","Pick a canonical form and document it, or your field names will not match anyone else's."),
("AI SBOM spec","Models has both a hash value AND a hash algorithm. Datasets has only 'Dataset hash' with no algorithm field.","An asymmetry in the spec itself. Decide whether your checker warns on it."),
("AI SBOM spec","Autonomy and agentic AI were CONSIDERED and DELIBERATELY EXCLUDED. The group said it may be handled through jurisdictional safety requirements instead.","Do not add an autonomy field and call it conformant. SPDX 3.1-dev adds isoAutomationLevel; that is SPDX's choice, not the G7's."),
("SBOM v2.1","17 fields, not 18. Split is 9 metadata + 8 component data.","Get the count right in your docs. Ten of the 17 are new since 2021."),
("SBOM v2.1","Uniformly 'should'. Zero instances of 'shall', zero of 'recommended', one 'must' and it is prose about sharing.","Same problem as the AI spec: no conformance tiers exist to extract."),
("SBOM v2.1","SBOM Timestamp requires RFC 9557, not RFC 3339. RFC 9557 allows bracketed suffixes such as [America/Chicago].","Both CycloneDX and SPDX validate timestamps as plain date-time. A strictly conformant timestamp can fail their schemas. Handle this explicitly."),
("SBOM v2.1","SWID Tags were removed from the accepted format list.","A 2026-conformant checker should no longer treat SWID as an SBOM format."),
("CycloneDX","The JSON Schema does NOT enforce that modelCard belongs to a machine-learning-model component, nor that data belongs to a data component. Those rules exist only in description strings.","A schema-only validator happily accepts a library with a model card attached. If you want to enforce the standard as written you implement this yourself. This is a genuine reason for your checker to exist."),
("CycloneDX","bom-ref uniqueness and refLinkType resolution are entirely unenforced. Every bom-ref-bearing object shares one flat namespace.","Build the index and resolve dangling refs yourself. Expect broken refs in real files."),
("CycloneDX","modelCard.modelParameters.datasets[] is a oneOf: inline data OR a ref, never both. An empty object {} validates. {'type':'dataset','ref':'x'} fails BOTH branches.","Classic parser trap. Test it explicitly in your test pack."),
("CycloneDX","performanceMetric.value and confidenceInterval bounds are STRINGS. energyMeasure.value and co2Measure.value are NUMBERS.","Do not let a code generator unify them."),
("CycloneDX","inputOutputMLParameters allows exactly ONE key: format. modelParameters has NO properties array.","The official cdx:ai-ml property taxonomy documentation contains examples that violate this and would fail validation. Put properties on modelCard.properties or component.properties."),
("CycloneDX","THE PYTHON AND JAVASCRIPT LIBRARIES DO NOT SUPPORT modelCard OR componentData. In cyclonedx-python-lib both are commented-out stubs marked TODO since CDX 1.5.","This is arguably the single biggest practical obstacle to AI BOM adoption, and it is in the two languages the ML world actually uses. Go, Java and .NET are complete. Consider this a product opportunity, not just a note."),
("CycloneDX","classification and sensitiveData are unconstrained free text. No vocabulary at all.","Layer your own policy if you want consistency."),
("SPDX","The External Property Restrictions (releaseTime, suppliedBy, downloadLocation, packageVersion, primaryPurpose all minCount 1 on AIPackage) are in the SPEC TEXT but NOT in the published SHACL or JSON Schema.","pyshacl and ajv will pass documents that violate the spec. Hard-code these five yourself. Same for DatasetPackage's five, where originatedBy is tightened to exactly 1."),
("SPDX","Profile conformance requires exactly one hasConcludedLicense and exactly one hasDeclaredLicense relationship per AIPackage and per DatasetPackage. Also not in SHACL.","Two more rules to hand-implement."),
("SPDX","datasetType is the ONLY property in either AI or Dataset profile with minCount 1 that is actually enforced.","Everything else is optional in the machine-readable artifacts."),
("SPDX","No baseModel, parentModel, derivedFrom or finetunedFrom exists in 3.0.1. Even 3.1-dev's finetunedOn points at a DATASET, not a parent model.","'Fine-tuned from Llama 3' has no interoperable representation. Use ancestorOf, or model the run as a Build with the parent as hasInput. No two producers will agree."),
("SPDX","Tag-value is NOT supported in SPDX 3.x. It is an RDF model: JSON-LD, Turtle, N-Triples, RDF/XML.","spdx/tools-golang still ships a tagvalue reader but targets 2.x only."),
("SPDX","spdx/tools-python is experimental, WRITE-ONLY for SPDX 3, and pinned to a 2023-era model commit. Its own README says do not use it in production.","Use spdx-python-model (generated from the ontology, so AI and Dataset profiles are covered) or generate bindings with shacl2code."),
("SPDX","safetyRiskAssessment uses the EU general product safety methodology from Regulation (EC) 765/2008 Article 20, and the spec explicitly warns it DIFFERS from the EU AI Act risk categorisation.","Anyone using it as an AI Act compliance field is misreading the spec. Worth saying out loud to customers."),
("Both formats","ZERO fields for adversarial robustness, prompt injection, jailbreak resistance, data poisoning, watermarking, guardrails or red-teaming. Verified by full-text scan of both schemas.","The AI-specific security content the G7 spec asks for has no interoperable home anywhere. This is the clearest unclaimed ground in the whole exercise."),
("Both formats","No open-weight vocabulary. No way to say open weight vs open architecture vs open data vs open training.","A named requirement of the spec with no representation in either format."),
("Both formats","No statistical characterisation of datasets. No operational KPIs. No context length. No modality in-schema.","Four more named requirements with no home."),
("FDA","Section 524B(b)(3) requires an SBOM by statute. The Feb 2026 guidance asks for NTIA 2021 baseline attributes plus per-component SUPPORT LEVEL and END-OF-SUPPORT DATE, plus known vulnerabilities including CISA KEV entries.","Support level and end-of-support date are FDA-specific and are NOT in the CRA, NOT in the 2026 minimum elements, and NOT native fields in either format. That is a concrete, sellable gap."),
("FDA","The SBOM text is WORD-FOR-WORD IDENTICAL across Sept 2023, June 2025 and Feb 2026. The 2026 change was a quality-system realignment to ISO 13485, not a cybersecurity change.","The evidence spec has been stable for about two and a half years. There is no pending FDA SBOM change to chase."),
("FDA","FDA names NO format. SPDX, CycloneDX and SWID appear nowhere in the guidance. It asks for machine-readable plus NTIA minimum elements and says industry-accepted formats are encouraged.","Format-agnostic output is a requirement, not a nicety."),
("FDA","The AI guidance (Jan 2025) is STILL DRAFT after 19 months and contains ZERO mentions of SBOM, bill of materials, datasheet, data card or provenance. Its model card is explicitly voluntary and explicitly not a template.","No regulator anywhere currently requires a machine-readable AI or model BOM. You are building ahead of the mandate, which is the opportunity and the risk in one sentence."),
("FDA","Failing to provide an SBOM is NOT a prohibited act under 21 USC 331(q)(3), which cites only 524B(b)(2). It is a submission-content deficiency enforced through Refuse to Accept, in force since 1 October 2023.","The enforcement path is RTA and review deficiency, not an enforcement action. Say it accurately."),
("EU CRA","Annex I Part II point 1 requires an SBOM 'in a commonly used and machine-readable format covering AT THE VERY LEAST the top-level dependencies'.","That is a FLOOR, not a ceiling, and it is shallower than FDA, which wants upstream transitive dependencies. Two different depth policies, one tool."),
("EU CRA","Article 13(24) reserves the SBOM format and elements to a Commission implementing act. NONE HAS BEEN ADOPTED. The CRA names no format.","A live, dated regulatory vacuum you can sell into."),
("EU CRA","The SBOM is NOT a mandatory user-facing deliverable (Annex II point 9 is conditional). It IS mandatory in the technical documentation (Annex VII) and on reasoned request from a market surveillance authority (Article 13).","Opposite of FDA, which requires it proactively in every submission and continuously to users."),
("EU CRA","Article 14 reporting applies from 11 SEPTEMBER 2026 and, per Article 69(3), reaches EVERY in-scope product already on the market. Clocks: 24 hours early warning, 72 hours notification, 14 days final report after a fix is available.","Two weeks out as of this writing, and retroactive to the installed base. The sharpest near-term sales narrative available."),
("EU CRA","Penalties for breaching Annex I or Articles 13/14 are up to EUR 15,000,000 or 2.5% of worldwide turnover. Supplying incorrect or incomplete information in reply to a request is separately up to EUR 5,000,000 or 1%.","The SBOM requirement sits in the TOP penalty tier. An inaccurate SBOM given under an Article 13 request is independently sanctionable."),
("EU CRA","No harmonised standards have been cited in the Official Journal, and ZERO notified bodies are designated as of 29 August 2026 despite Chapter IV applying since 11 June 2026.","Annex III Class I manufacturers legally cannot self-assess while no standards exist (Article 32(2)) and there is nobody to assess them. A structural bottleneck with full application on 11 December 2027."),
("EU AI Act","Annex IV points 2(a), 2(c) and 2(d) are the closest thing in EU law to a mandatory AI bill of materials: third-party pre-trained models and tools, system architecture and compute, and dataset datasheets with provenance.","But it is NARRATIVE technical documentation with no machine-readable requirement. Article 12 of the CRA bridges them: a high-risk AI system claiming the AI Act Article 15 cybersecurity presumption must satisfy the CRA SBOM requirement."),
]
for idx, row in enumerate(T, start=1):
    ws4.append([idx] + list(row))
    r = ws4.max_row
    for c in range(1, len(H4) + 1):
        cell = ws4.cell(row=r, column=c)
        cell.font = Font(name=FONT, size=9)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = BORD
    ws4.cell(row=r, column=1).alignment = Alignment(vertical="top", horizontal="center")
    ws4.cell(row=r, column=2).fill = CLUSTER_FILL
    ws4.cell(row=r, column=2).font = Font(name=FONT, size=9, bold=True)
finish(ws4, [4, 16, 86, 86], len(H4))

# ============================================================
# SHEET 5 - sources
# ============================================================
ws5 = wb.create_sheet("Sources")
H5 = ["Document", "Date", "Where to get it", "Note"]
ws5.append(H5)
style_header(ws5, len(H5))
SRC = [
("Software Bill of Materials for AI - Minimum Elements (CISA + G7)","Published 12 May 2026 (BSI); CISA announced 13 May 2026","https://www.bsi.bund.de/SharedDocs/Downloads/EN/BSI/KI/SBOM-for-AI_minimum-elements.pdf","cisa.gov returns 403 to automated fetching. Germany's BSI is a co-publisher and hosts the identical original. The document itself carries NO date beyond '2026' and no version number. The widely repeated '16 June 2026' is a law firm alert date, not the publication date."),
("2026 Minimum Elements for a Software Bill of Materials (SBOM), v2.1","29 July 2026","https://www.ic3.gov/CSA/2026/260729.pdf","Also at media.defense.gov, which 403s. 18 authoring agencies: CISA, NSA, FBI plus 15 international. TLP:CLEAR. 23 pages. Note: NCSC UK is NOT an author of this one."),
("A shared G7 vision on Software Bill of Materials for AI","19 May 2025","https://www.acn.gov.it/portale/documents/d/guest/paper_sbom-for-ai_19may2025_-clean-2","The predecessor document that defined the original clusters. Some were added, adjusted and removed on the way to the 2026 minimum elements."),
("CycloneDX 1.7 JSON Schema","Released 21 October 2025","https://raw.githubusercontent.com/CycloneDX/specification/master/schema/bom-1.7.schema.json","ECMA-424 2nd Edition. The ML-BOM object graph is byte-for-byte identical between 1.6 and 1.7, so a 1.6 parser needs zero structural changes. modelCard arrived in 1.5; environmentalConsiderations in 1.6."),
("CycloneDX Authoritative Guide to AI/ML-BOM","1st edition rev 1, 10 June 2026","https://cyclonedx.org/guides/OWASP_CycloneDX-Authoritative-Guide-to-AI-ML-BOM-en.pdf","Maps to EU AI Act Article 53 and Annex XI. Its own examples contain schema-invalid property placements."),
("SPDX 3.0.1 model, AI and Dataset profiles","17 December 2024","https://spdx.github.io/spdx-spec/v3.0.1/model/AI/AI/","Normative machine-readable model at https://spdx.org/rdf/3.0.1/spdx-model.ttl. 3.1 is RC1 only as of August 2026. ISO/IEC 5962:2021 still covers version 2.2.1, not 3.x."),
("FD&C Act section 524B / 21 U.S.C. 360n-2","Effective 29 March 2023","https://www.govinfo.gov/content/pkg/USCODE-2024-title21/html/USCODE-2024-title21-chap9-subchapV-partA-sec360n-2.htm","(b)(3) is the SBOM requirement. 'Cyber device' is a three-part conjunctive test."),
("FDA Cybersecurity in Medical Devices (QMS considerations and premarket content)","3 February 2026, final","https://www.fda.gov/media/119933/download","Docket FDA-2021-D-1158. Supersedes the June 2025 version. SBOM content is unchanged since September 2023."),
("FDA Refuse to Accept policy for cyber devices","30 March 2023; enforced from 1 October 2023","https://www.govinfo.gov/content/pkg/FR-2023-03-30/html/2023-06646.htm","88 FR 19148. The grace window was temporal, not substantive."),
("FDA AI-Enabled Device Software Functions (draft)","7 January 2025, STILL DRAFT","https://www.fda.gov/media/184856/download","Docket FDA-2024-D-4488. Zero mentions of SBOM. Model card is voluntary and explicitly not a template."),
("FDA AI-Enabled Medical Devices list","Page updated 16 June 2026; data through 30 March 2026","https://www.fda.gov/media/178541/download?attachment","1,524 devices. 76% radiology. 333 authorised in 2025 alone. FDA states it is not comprehensive."),
("FDA Establishment Registration and Device Listing bulk files","Updated every Sunday","https://www.fda.gov/medical-devices/device-registration-and-listing/establishment-registration-and-medical-device-listing-files-download","13 zipped pipe-delimited files. Join registration_listing product codes against the AI device CSV's Primary Product Code to build the prospect list."),
("Regulation (EU) 2024/2847 (Cyber Resilience Act)","In force 10 December 2024","http://publications.europa.eu/resource/celex/32024R2847","eur-lex is WAF-blocked to scripted access; the Publications Office cellar service serves the same authentic text. Annex I Part II point 1 is the SBOM requirement."),
("Regulation (EU) 2024/1689 (AI Act)","-","http://publications.europa.eu/resource/celex/32024R1689","Annex IV points 2(a), (c), (d) are the AI-BOM-adjacent technical documentation elements. Narrative, not machine-readable."),
("NANDO / Single Market Compliance Space","Checked 29 August 2026","https://webgate.ec.europa.eu/single-market-compliance-space/","CRA legislation record exists (id 167953) with zero notified body designations and zero harmonised standards attached."),
]
for row in SRC:
    ws5.append(list(row))
    r = ws5.max_row
    for c in range(1, len(H5) + 1):
        cell = ws5.cell(row=r, column=c)
        cell.font = Font(name=FONT, size=9)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = BORD
    ws5.cell(row=r, column=1).font = Font(name=FONT, size=9, bold=True)
finish(ws5, [46, 22, 62, 86], len(H5))

# ============================================================
# SHEET 6 - legend
# ============================================================
ws6 = wb.create_sheet("Read me")
ws6["A1"] = "AI SBOM crosswalk"
ws6["A1"].font = Font(name=FONT, size=16, bold=True, color="1F3864")
notes = [
 "",
 "What this is",
 "A field-level mapping of the CISA and G7 'Software Bill of Materials for AI - Minimum Elements' onto the two formats the",
 "industry actually uses, CycloneDX 1.7 and SPDX 3.0.1, with the gaps marked in both directions.",
 "",
 "Why it exists",
 "The G7 document names no format. The string 'CycloneDX' does not appear in it at all, and 'SPDX' appears exactly once,",
 "inside an example. So there is no official crosswalk, and as far as could be found, no published unofficial one either.",
 "",
 "Sheets",
 "  AI SBOM crosswalk     All 50 elements across the 7 clusters, mapped to both formats.",
 "  SBOM v2.1 - 17 fields The general software SBOM baseline the AI elements sit ON TOP OF, not instead of.",
 "  Practices (6)         The six practices and processes from the 2026 minimum elements.",
 "  Traps and gaps        Everything that will bite an implementer, including the regulatory asymmetries.",
 "  Sources               Every primary document, where to actually get it, and the access notes that matter.",
 "",
 "Fit ratings",
 "  Direct        The format has a purpose-built field that carries this element.",
 "  Partial       A field exists but is narrower, differently scoped, or uses an incompatible vocabulary.",
 "  Properties    Only expressible through a free-form property bag or an external link. Not interoperable.",
 "  None          No representation at all.",
 "",
 "The one thing to take away",
 "Both formats are strong on identity, licensing, dependencies and vulnerabilities. Both are weak on model lineage.",
 "And neither has a single field for adversarial robustness, prompt injection, jailbreak resistance, data poisoning,",
 "watermarking, guardrails or red-teaming. That was verified by full-text scan of both schemas: zero hits, every term.",
 "The AI-specific security content the G7 asked for has no interoperable home anywhere.",
 "",
 "Compiled 29 August 2026 from the primary documents. Verify against the sources before relying on any of it.",
]
for i, line in enumerate(notes, start=2):
    ws6.cell(row=i, column=1, value=line)
    f = Font(name=FONT, size=10)
    if line in ("What this is","Why it exists","Sheets","Fit ratings","The one thing to take away"):
        f = Font(name=FONT, size=11, bold=True, color="1F3864")
    ws6.cell(row=i, column=1).font = f
ws6.column_dimensions["A"].width = 125
ws6.sheet_view.showGridLines = False

wb.move_sheet("Read me", offset=-5)
wb.save(os.path.join(_CW, "AI-SBOM-crosswalk.xlsx"))
print("saved")
